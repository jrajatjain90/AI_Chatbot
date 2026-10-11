import os, re, json
from datetime import date
from decimal import Decimal
from fastapi.responses import FileResponse
from uuid import uuid4
import pandas as pd
from fastapi import FastAPI
from pydantic import BaseModel
from databricks import sql
from databricks.sdk.core import Config
from openai import OpenAI
from io import BytesIO
from fastapi import Response
from openpyxl import Workbook
from openpyxl.chart import BarChart, LineChart, PieChart, Reference
from typing import Any
import traceback
from collections import defaultdict


app = FastAPI(title="Sales Chatbot API")
cfg = Config()

WAREHOUSE_ID = os.getenv("WAREHOUSE_ID", "da19d04f9077024b")
MODEL = "databricks-meta-llama-3-3-70b-instruct"

# ---------- connections ----------
def run_query(query):
    with sql.connect(
        server_hostname=cfg.host,
        http_path=f"/sql/1.0/warehouses/{WAREHOUSE_ID}",
        credentials_provider=lambda: cfg.authenticate,
    ) as conn:
        with conn.cursor() as cursor:
            cursor.execute(query)
            rows = cursor.fetchall()
            cols = [c[0] for c in cursor.description]
    df = pd.DataFrame(rows, columns=cols)
    for c in df.columns:
        if df[c].map(lambda v: isinstance(v, Decimal)).any():
            df[c] = df[c].astype(float)
    return df

def insert_chat_log(request_id, question,sql_query):
     with sql.connect(
        server_hostname=cfg.host,
        http_path=f"/sql/1.0/warehouses/{WAREHOUSE_ID}",
        credentials_provider=lambda: cfg.authenticate,
    ) as conn:
         with conn.cursor() as cursor:
            cursor.execute("""INSERT INTO chat_bot.gold.chatbot_query_logs
                (request_id, question, sql_query, created_at)
                VALUES (?, ?, ?, current_timestamp())""", ( sql_query,question,request_id))        
            
def insert_feedback(request_id,feedback):
    with sql.connect(
        server_hostname=cfg.host,
        http_path=f"/sql/1.0/warehouses/{WAREHOUSE_ID}",
        credentials_provider=lambda: cfg.authenticate,
    ) as conn:
        with conn.cursor() as cursor:
            cursor.execute(
                """
                INSERT INTO chat_bot.gold.chatbot_feedback
                (request_id, feedback, created_at)
                VALUES (?, ?, current_timestamp())
                """,(request_id,feedback)
            )

            
def get_llm():
    token = cfg.authenticate()["Authorization"].replace("Bearer ", "")
    return OpenAI(api_key=token, base_url=f"{cfg.host}/serving-endpoints")

def ask_llm(system, user, max_tokens=300):
    resp = get_llm().chat.completions.create(
        model=MODEL,
        messages=[{"role": "system", "content": system}, {"role": "user", "content": user}],
        temperature=0,
        max_tokens=max_tokens,
    )
    return resp.choices[0].message.content

# ---------- helpers ----------
def parse_json(text):
    text = text.strip()
    if text.startswith("```"):
        text = text.strip("`")
        if text.lower().startswith("json"):
            text = text[4:]
    start, end = text.find("{"), text.rfind("}")
    return json.loads(text[start:end + 1])

def clean_sql(text):
    text = text.strip()
    if text.startswith("```"):
        text = text.strip("`")
        if text.lower().startswith("sql"):
            text = text[3:]
    return text.strip().rstrip(";")

# ---------- SQL generation ----------
def build_system():
    return f"""You are a Spark SQL expert. Convert the user's question into one SQL query.
Today's date is {date.today().isoformat()}. Use it for words like "last month" or "this year".

Tables:
chat_bot.gold.sales_enriched (one row per sale)
  sale_id, sale_date (date), sale_month (string yyyy-MM), customer_id, customer_name,
  region, segment, product_id, product_name, category,
  price (unit price), quantity, total_amount (revenue of the sale)

chat_bot.gold.monthly_sales_by_category
  sale_month (string yyyy-MM), category, revenue, units_sold, orders

Allowed values:
- region: North, South, East, West
- segment: Retail, Corporate, Online
- category: Laptops, Phones, Accessories, Audio
- Data covers 2026-04 to 2026-09

Business definitions:
- revenue = SUM(total_amount)
- units sold = SUM(quantity)
- orders = COUNT(DISTINCT sale_id)
- average order value = AVG(total_amount)
- customers = COUNT(DISTINCT customer_id)

Rules:
- Use only the tables and columns above, with full names like chat_bot.gold.sales_enriched
- Only SELECT queries
- Use sales_enriched unless the question is only about monthly totals by category
- Use a LIMIT only for "top N" questions, otherwise add LIMIT 50 when the query returns many rows
- For monthly or trend questions, order by sale_month
- Give columns clear aliases like revenue, units, orders
- Return only the SQL, with no explanation and no markdown

Examples:
Q: Revenue by category
SQL: SELECT category, SUM(total_amount) AS revenue FROM chat_bot.gold.sales_enriched GROUP BY category ORDER BY revenue DESC

Q: Top 3 customers by revenue
SQL: SELECT customer_name, SUM(total_amount) AS revenue FROM chat_bot.gold.sales_enriched GROUP BY customer_name ORDER BY revenue DESC LIMIT 3

Q: Monthly orders for Phones
SQL: SELECT sale_month, COUNT(DISTINCT sale_id) AS orders FROM chat_bot.gold.sales_enriched WHERE category = 'Phones' GROUP BY sale_month ORDER BY sale_month"""

def generate_sql(question):
    return clean_sql(ask_llm(build_system(), question))

BLOCKED = ["insert", "update", "delete", "drop", "alter", "create", "truncate", "merge", "grant"]

def is_safe(q):
    s = q.strip().lower()
    if not (s.startswith("select") or s.startswith("with")) or ";" in s:
        return False
    return not any(re.search(rf"\b{w}\b", s) for w in BLOCKED)

def tables_ok(q):
    refs = re.findall(r"(?:from|join)\s+([\w.]+)", q, flags=re.I)
    return all(r.lower().startswith("chat_bot.gold.") for r in refs)

def run_sql(q):
    if not (is_safe(q) and tables_ok(q)):
        raise ValueError(f"Blocked unsafe SQL: {q}")
    return run_query(q)

def fix_sql(question, bad_sql, error):
    return clean_sql(ask_llm(build_system(),
        f"Question: {question}\n\nThis SQL failed:\n{bad_sql}\n\nError:\n{error}\n\nReturn a corrected SQL query only."))

def ask_sql(question, max_retries=2):
    q = generate_sql(question)
    
    for attempt in range(max_retries + 1):
        try:
            return q, run_sql(q)
        except Exception as e:
            if attempt == max_retries:
                return q, f"FAILED: {e}"
            q = fix_sql(question, q, str(e)[:500])

def explain_Sql(question,result):
    metrics,largest_increase,largest_decrease=calculate_revenue_metrics(result)
   
    build = """
            You are an expert business data analyst for a business intelligence chatbot.

            Your task is to explain SQL results and precomputed metrics clearly and accurately.

            Return one valid JSON object only. Do not use Markdown code fences or
            include text outside the JSON object.

            Required JSON structure:
            {
            "key_findings": [
                {
                "title": "Short finding title",
                "value": "Verified metric or percentage",
                "period": "Relevant period",
                "type": "positive",
                "description": "Evidence-based explanation"
                }
            ],
            "trend_summary": "Summary of the overall trend.",
            "recommendation": "Evidence-based next step."
            }

            Rules:
            1. Use only the supplied SQL results and calculated metrics.
            2. Treat Python-calculated metrics as the source of truth.
            3. Do not recalculate, modify, round differently, or contradict
            any supplied calculated metric.
            4. Preserve the exact values, region names, and periods from
            the calculated metrics.
            5. Do not invent figures, causes, or comparisons.
            6. Recommendations must be supported by the available evidence.
            7. If the data is insufficient, state that clearly.
            8. Return at most 3 key findings.
            9. Each finding must contain exactly these fields:
            title, value, period, type, description.
            10. The type must be "positive", "negative", or "neutral".
            11. Use exactly these top-level keys:
                key_findings, trend_summary, recommendation.
            12. All JSON strings must use double quotes.
            13. Do not include trailing commas.
            14. Return valid JSON only.
"""




    user_message=(
      f"user question:{question}\n"
      f"Query results: {json.dumps(result, default=str)}\n"
      f"calculated metrics:{metrics}\n"
      f"largest_increase:{largest_increase}\n"
      f"largest_decrease:{largest_decrease}\n"
      "Use the calculated metric as source of truth"
      )
    
    resp=ask_llm(
        build,
        user_message
    )
    try:
        return json.loads(resp)
    except json.JSONDecodeError as exc:
        raise ValueError("The LLM returned invalid json for the insights response") from exc

def calculate_revenue_metrics(results):
    monthly_region = defaultdict(dict)

    for row in results:
        month = row["sale_month"]
        region = row["region"]
        revenue = float(row["revenue"])
        monthly_region[region][month] = revenue

    metrics = {
        "region_averages": {},
        "monthly_changes": []
        }

    for region, months in monthly_region.items():
        ordered = sorted(months.items())

        metrics["region_averages"][region] = round(
            sum(value for _, value in ordered) / len(ordered), 2
        )

        for i in range(1, len(ordered)):
            previous_month, previous_value = ordered[i - 1]
            current_month, current_value = ordered[i]

            if previous_value != 0:
                change_pct = (
                    (current_value - previous_value)
                    / previous_value
                ) * 100

                metrics["monthly_changes"].append({
                    "region": region,
                    "from_month": previous_month,
                    "to_month": current_month,
                    "change_pct": round(change_pct, 2)
                })
    changes= metrics['monthly_changes']                    
    largest_increase=max(
        changes, key=lambda x:x['change_pct']        
    )

    largest_decrease= min(
        changes, key=lambda x:x['change_pct']
    )

    return (metrics, largest_increase,largest_decrease)




# ---------- answer + chart ----------
def summarize(question, df):
    return ask_llm(
        "You answer business questions using only the data provided. Be brief and clear. "
        "Do not invent numbers. If the data is empty, say no data was found.",
        f"Question: {question}\n\nData (CSV):\n{df.head(20).to_csv(index=False)}", 200)

CHART_SYSTEM = """You choose a chart for a SQL query result.
Return only JSON with keys: chart_type, x, y, series, title.
- chart_type: "bar", "line", "pie" or "none"
- "line" when x is a month or date, "bar" for categories, "pie" only for shares of a total with at most 6 categories, "none" for a single value or text-only results
- x: the text or date column for the labels. y: the numeric column. series: another text column that splits the data into groups, or null
- x, y and series must be copied exactly from the column names provided in the user message. Never use generic words like value, category, sales or label
- title: short, plain English"""

def suggest_chart(question, df):
    none = {"chart_type": "none"}
    if len(df) < 2:
        return none
    info = (f"Question: {question}\nColumns and types: {df.dtypes.astype(str).to_dict()}\n"
            f"Sample rows:\n{df.head(5).to_csv(index=False)}")
    try:
        spec = parse_json(ask_llm(CHART_SYSTEM, info, 150))
    except Exception:
        return none
    num_cols = df.select_dtypes("number").columns.tolist()
    txt_cols = [c for c in df.columns if c not in num_cols]
    if spec.get("y") not in df.columns and num_cols:
        spec["y"] = num_cols[0]
    if spec.get("x") not in df.columns and txt_cols:
        spec["x"] = txt_cols[0]
    ok = (spec.get("chart_type") in ("bar", "line", "pie")
          and spec.get("x") in df.columns and spec.get("y") in df.columns
          and pd.to_numeric(df[spec["y"]], errors="coerce").notna().all())
    if not ok:
        return none
    if spec.get("series") not in df.columns:
        spec["series"] = None
    return spec

def chat_sql(question):
    q, result = ask_sql(question)

    if isinstance(result, str):
        return {"answer": "Sorry, I could not answer that from the sales data.",
                "sql": q, "rows": [], "chart": {"chart_type": "none"}}
    rows = json.loads(result.head(50).to_json(orient="records", date_format="iso"))
    return {"answer": summarize(question, result), "sql": q, "rows": rows,
            "chart": suggest_chart(question, result)}

def chat_suggestions(question):
    build = f"""
    Act as a question suggestion provider.

    Based on the user's question: "{question}"

    Generate exactly 3 relevant follow-up questions.
    Return only a Python list of 3 strings.
    """

    response = ask_llm(build,question)
    return response




# ---------- API ----------
class ChatRequest(BaseModel):
    question: str


@app.get("/")
def root():
    return FileResponse(os.path.join(os.path.dirname(__file__), "index.html"))

@app.get("/health")
def health():
    return {"status": "ok"}

@app.post("/chat")
def chat(req: ChatRequest):

    try:
        request_id = str(uuid4())
        response= chat_sql(req.question)
        insert_chat_log(response['sql'],req.question,request_id) 
        return {
            **response,"request_id":request_id
            }
    except Exception as e:
        return {
            "request_id":request_id if "request_id" in locals() else "",
            "answer": "Something went wrong.", "error": f"{type(e).__name__}: {e}",
                "sql": "", "rows": [], "chart": {"chart_type": "none"}
                }
        

class SuggestionRequest(BaseModel):
    question:str

@app.post("/suggestions")
def get_suggestions(req:SuggestionRequest):
    return chat_suggestions(req.question)            
        
class FeedbackRequest(BaseModel):
    request_id:str
    feedback:str

class SqlInsights(BaseModel):
    question:str
    results:list[dict[str,Any]]

@app.post("/insights")
def get_sql_insights(data:SqlInsights):
    try:
        return explain_Sql(data.question,data.results)
    except Exception as e:
        traceback.print_exc()
        return {"error": str(e)}
    

@app.post("/feedback")
def submit_feedback(data:FeedbackRequest):
    
    insert_feedback(
        data.request_id,
        data.feedback
    )

    return {"status":"ok"}

class ExportRequest(BaseModel):
    rows: list
    chart: dict = {}

@app.post("/export")
def export(req: ExportRequest):
    df = pd.DataFrame(req.rows)
    wb = Workbook()
    ws = wb.active
    ws.title = "Data"
    ws.append(list(df.columns))
    for r in df.itertuples(index=False):
        ws.append(list(r))
    s = req.chart or {}
    kind = s.get("chart_type")
    if kind in ("bar", "line", "pie") and s.get("x") in df.columns and s.get("y") in df.columns:
        ch = {"bar": BarChart, "line": LineChart, "pie": PieChart}[kind]()
        ch.title = s.get("title", "")
        ch.width, ch.height = 22, 11
        if s.get("series") in df.columns and kind != "pie":
            p = df.pivot_table(index=s["x"], columns=s["series"], values=s["y"], aggfunc="sum").reset_index()
            src = wb.create_sheet("Chart data")
            src.append([str(c) for c in p.columns])
            for r in p.itertuples(index=False):
                src.append(list(r))
            data = Reference(src, min_col=2, max_col=len(p.columns), min_row=1, max_row=len(p) + 1)
            cats = Reference(src, min_col=1, min_row=2, max_row=len(p) + 1)
        else:
            cols = list(df.columns)
            data = Reference(ws, min_col=cols.index(s["y"]) + 1, min_row=1, max_row=len(df) + 1)
            cats = Reference(ws, min_col=cols.index(s["x"]) + 1, min_row=2, max_row=len(df) + 1)
        ch.add_data(data, titles_from_data=True)
        ch.set_categories(cats)
        wb.create_sheet("Chart").add_chart(ch, "A1")
    buf = BytesIO()
    wb.save(buf)
    return Response(buf.getvalue(),
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": 'attachment; filename="sales_export.xlsx"'})






    
    
        
    
    


